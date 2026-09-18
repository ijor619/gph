import docx
from docx.shared import Pt
import datetime
import calendar

MONTHS_RU = {
    1: "января", 2: "февраля", 3: "марта", 4: "апреля",
    5: "мая", 6: "июня", 7: "июля", 8: "августа",
    9: "сентября", 10: "октября", 11: "ноября", 12: "декабря"
}

def get_russian_date(dt: datetime.date) -> str:
    day = f"«{dt.day:02d}»" if dt.day < 10 else f"«{dt.day}»"
    return f"{day} {MONTHS_RU[dt.month]} {dt.year}г."

def get_russian_date_pd(dt: datetime.date) -> str:
    day = f"«{dt.day:02d}»" if dt.day < 10 else f"«{dt.day}»"
    return f"{day} {MONTHS_RU[dt.month]} {dt.year} г."

def add_six_months(source_date: datetime.date) -> datetime.date:
    month = source_date.month - 1 + 6
    year = source_date.year + month // 12
    month = month % 12 + 1
    max_days = calendar.monthrange(year, month)[1]
    day = min(source_date.day, max_days)
    return datetime.date(year, month, day)

def make_fio_initials(fio_str: str) -> str:
    """
    Формирует фамилию с инициалами.
    'Абдунабиев Садамбек Зафарович' -> 'Абдунабиев С.З.'
    'Муртузалиев Зия Азер оглу' -> 'Муртузалиев З.А.О.'
    """
    parts = fio_str.strip().split()
    if not parts:
        return ""
    last_name = parts[0].capitalize()
    initials = "".join([f"{p[0].upper()}." for p in parts[1:]])
    return f"{last_name} {initials}".strip()

def fill_gpd_contract(
    template_path: str,
    output_path: str,
    data: dict
) -> str:
    """Заполняет договор ГПХ по шаблону с сохранением форматирования"""
    doc = docx.Document(template_path)
    
    # Дата договора: если не задана вручную, по правилу берём СЕГОДНЯ + 1 ДЕНЬ К КАЛЕНДАРЮ
    if isinstance(data.get("contract_date"), datetime.date):
        dt_start = data["contract_date"]
    else:
        dt_start = datetime.date.today() + datetime.timedelta(days=1)
    
    # Срок действия: ровно +6 месяцев от даты подписания
    dt_end = add_six_months(dt_start)
    start_date_str = get_russian_date(dt_start)
    end_date_str = get_russian_date(dt_end)
    
    contract_num = str(data.get("contract_num", "1"))
    fio = data.get("fio", "").strip()
    fio_short = make_fio_initials(fio)
    citizenship = data.get("citizenship", "").strip()
    birth_date = data.get("birth_date", "").strip()
    birth_place = data.get("birth_place", "").strip()
    passport_str = data.get("passport_str", "").strip()
    work_doc_full = data.get("work_doc_full", "").strip()
    work_doc_table = data.get("work_doc_table", f"Основание: {work_doc_full}").strip()
    reg_address = data.get("reg_address", "").strip()
    stay_basis = data.get("stay_basis", work_doc_table).strip()
    stay_issuer = data.get("stay_issuer", "").strip()
    inn = data.get("inn", "(заполнить)").strip()
    snils = data.get("snils", "(заполнить)").strip()
    bik = data.get("bik", "(заполнить)").strip()
    rs = data.get("rs", "(заполнить)").strip()
    ks = data.get("ks", "(заполнить)").strip()

    # 1. Заголовок (номер договора)
    doc.paragraphs[0].text = f"Договор гражданско-правового характера № {contract_num}"
    doc.paragraphs[0].runs[0].bold = True
    doc.paragraphs[0].runs[0].font.name = "Times New Roman"
    doc.paragraphs[0].runs[0].font.size = Pt(11)

    # 2. Дата договора в шапке (завтрашний день: +1 день к календарю)
    doc.paragraphs[2].text = f"г. Санкт-Петербург\t{start_date_str}"
    for r in doc.paragraphs[2].runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)

    # 3. Преамбула (Исполнитель)
    doc.paragraphs[4].text = (
        f"Гражданин {citizenship} {fio} {birth_date} года рождения, "
        f"место рождения: {birth_place}, документ удостоверяющий личность: {passport_str}; "
        f"основание для ведения трудовой деятельности: {work_doc_full}, "
        f"зарегистрированный по адресу: {reg_address}, "
        f"именуемый в дальнейшем «Исполнитель», с другой стороны, а совместно именуемые Стороны, "
        f"заключили настоящий Договор о нижеследующем:"
    )
    for r in doc.paragraphs[4].runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)

    # 4. Подтверждение пребывания (п. 5)
    doc.paragraphs[5].text = (
        f"Исполнитель подтверждает, что является гражданином {citizenship}, "
        f"пребывающим в РФ на основании {stay_basis} "
        f"выданного {stay_issuer}."
    )
    for r in doc.paragraphs[5].runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)

    # 5. Срок выполнения работ (п. 1.5)
    for p in doc.paragraphs:
        if "1.5. Срок выполнения работ" in p.text:
            p.text = (
                f"1.5. Срок выполнения работ по настоящему Договору с {start_date_str} "
                f"по {end_date_str} (с даты подписания 6 месяцев)"
            )
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(11)
            break

    # 6. Таблица реквизитов (Исполнитель)
    table = doc.tables[0]
    cell1 = table.rows[1].cells[1]
    cell1.text = ""
    lines = [
        "",
        f"{fio} ",
        "(Ф.И.О.)",
        "",
        f"{work_doc_table}",
        "",
        f"Адрес регистрации: {reg_address}",
        "",
        f"ИНН: {inn}",
        "",
        f"СНИЛС: {snils}",
        "",
        "Банковские реквизиты для перечисления оплаты: ",
        f"БИК: {bik}",
        f"Р/С: {rs}",
        f"К/С: {ks}"
    ]
    p_first = cell1.paragraphs[0]
    p_first.text = lines[0]
    for line in lines[1:]:
        cell1.add_paragraph(line)
    for p in cell1.paragraphs:
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)

    # 7. Таблица подписей
    cell2 = table.rows[2].cells[1]
    cell2.text = ""
    cell2.paragraphs[0].text = ""
    cell2.add_paragraph("")
    cell2.add_paragraph(f"_____________ / {fio_short}/ ")
    cell2.add_paragraph("    (подпись)              (Ф.И.О.)")
    for p in cell2.paragraphs:
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)

    doc.save(output_path)
    return output_path


def fill_pd_consent(
    template_path: str,
    output_path: str,
    data: dict
) -> str:
    """Заполняет согласие на обработку персональных данных"""
    doc = docx.Document(template_path)
    
    fio = data.get("fio", "").strip()
    fio_short = make_fio_initials(fio)
    passport_str = data.get("passport_str", "").strip()
    
    # Нормализация строки паспорта
    if not passport_str.lower().startswith("паспорт"):
        passport_str = f"паспорт: {passport_str}"
    else:
        passport_str = passport_str.replace("паспорт ", "паспорт: ")
        
    if not passport_str.endswith("г.") and not passport_str.endswith("г"):
        passport_str = f"{passport_str} г."
        
    reg_address = data.get("reg_address", "").strip()
    
    # Дата согласия: та же дата, что и в договоре (+1 день к календарю)
    if isinstance(data.get("contract_date"), datetime.date):
        dt = data["contract_date"]
    else:
        dt = datetime.date.today() + datetime.timedelta(days=1)
        
    date_str = get_russian_date_pd(dt)
    
    # 1. P2: Я, ФИО, паспорт: ... выдан ...
    doc.paragraphs[2].text = f"Я, {fio}, {passport_str} "
    
    # 2. P4: зарегистрированной(го) по адресу: ...
    doc.paragraphs[4].text = f"зарегистрированной(го) по адресу: {reg_address}"
    
    # 3. P62: строка подписи с инициалами и датой подписания
    sig_p = None
    for p in reversed(doc.paragraphs):
        if "________________________" in p.text or "/________" in p.text or "«" in p.text:
            sig_p = p
            break
            
    if sig_p:
        sig_p.text = f"{fio_short}  /________________________/                                                       {date_str}"
        for r in sig_p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(10)
            
    doc.save(output_path)
    return output_path
